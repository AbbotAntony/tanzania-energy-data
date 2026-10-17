"""Tests for the labelling sample, blind sheet and label validation."""

from __future__ import annotations

import pandas as pd
import pytest

from energy_nlp import labels as lb


def _corpus():
    rows = []
    pid = 1
    for year, n_gas, n_other in [(2022, 25, 12), (2023, 40, 30), (2024, 22, 9)]:
        for gas, n in [(True, n_gas), (False, n_other)]:
            for _ in range(n):
                rows.append((pid, pd.Timestamp(f"{year}-06-01"), "natural_gas;tpdc" if gas else "jnhpp",
                             f"https://x/{pid}", f"title {pid}", f"text {pid}"))
                pid += 1
    return pd.DataFrame(rows, columns=["post_id", "pub_date", "topics", "url", "title", "text"])


def test_draw_sample_strata_and_weights():
    c = _corpus()
    s = lb.draw_sample(c, seed=42)
    counts = s.groupby("stratum").size().to_dict()
    assert counts == {"2022_gas": 20, "2022_other": 10, "2023_gas": 20, "2023_other": 10,
                      "2024_gas": 20, "2024_other": 9}            # 2024_other has only 9
    # weights re-create the population size of every stratum
    assert s.groupby("stratum")["weight"].sum().round(6).to_dict() == \
        {"2022_gas": 25, "2022_other": 12, "2023_gas": 40, "2023_other": 30, "2024_gas": 22, "2024_other": 9}
    assert s["post_id"].is_unique
    assert s["sample_id"].tolist() == [f"S{i:03d}" for i in range(1, len(s) + 1)]


def test_draw_sample_reproducible_and_shuffled():
    c = _corpus()
    a, b = lb.draw_sample(c, seed=42), lb.draw_sample(c, seed=42)
    pd.testing.assert_frame_equal(a, b)
    assert not a["stratum"].is_monotonic_increasing          # order does not reveal strata
    assert set(lb.draw_sample(c, seed=7)["post_id"]) != set(a["post_id"])


def test_blind_sheet_hides_identifying_columns():
    sheet = lb.blind_sheet(lb.draw_sample(_corpus()))
    assert list(sheet.columns) == ["sample_id", "title", "text", "gas_mention", "gas_tone", "frame",
                                   "unsure", "notes"]
    assert (sheet[["gas_mention", "gas_tone", "frame"]] == "").all().all()


def test_draw_recode_subset_and_reproducible():
    design = lb.design_table(lb.draw_sample(_corpus()))
    r = lb.draw_recode(design, n=30, seed=43)
    assert len(set(r)) == 30 and set(r) <= set(design["sample_id"])
    assert r == lb.draw_recode(design, n=30, seed=43)


def test_write_sheet_refuses_to_overwrite(tmp_path):
    sheet = lb.blind_sheet(lb.draw_sample(_corpus()))
    path = tmp_path / "sheet.csv"
    lb.write_sheet(sheet, path)
    with pytest.raises(FileExistsError):
        lb.write_sheet(sheet, path)
    lb.write_sheet(sheet, path, force=True)


@pytest.mark.parametrize("encoding", ["utf-8-sig", "cp1252"])
def test_read_sheet_normalises_excel_input(tmp_path, encoding):
    df = pd.DataFrame({
        "sample_id": ["S001", "S002"], "title": ["Gas’s rise", "b"], "text": ["t", "u"],
        "Gas_Mention": ["Substantive ", "resource list"], "gas_tone": ["POSITIVE", "not_about_gas"],
        "frame": ["economic_opportunity", "resource_wealth"], "unsure": ["x", ""], "notes": ["", " hm "],
    })
    path = tmp_path / "done.csv"
    df.to_csv(path, index=False, encoding=encoding)
    out = lb.read_sheet(path)
    assert out["gas_mention"].tolist() == ["substantive", "resource_list"]
    assert out["gas_tone"].tolist() == ["positive", "not_about_gas"]
    assert out["unsure"].tolist() == [True, False]
    assert out["notes"].tolist() == ["", "hm"]
    assert "title" not in out.columns                            # no text in the label file


def _labels(rows):
    return pd.DataFrame(rows, columns=["sample_id", "gas_mention", "gas_tone", "frame"]).assign(
        unsure=False, notes="")


def test_validate_clean_labels_pass():
    ok = _labels([("S1", "substantive", "positive", "economic_opportunity"),
                  ("S2", "resource_list", "not_about_gas", "resource_wealth"),
                  ("S3", "energy_mix_list", "not_about_gas", "supply_reliability"),
                  ("S4", "none", "not_about_gas", "other")])
    assert lb.validate(ok).empty


def test_validate_flags_each_rule():
    bad = _labels([("S1", "", "positive", "economic_opportunity"),        # missing
                   ("S2", "substantive", "great", "economic_opportunity"),  # invalid code
                   ("S3", "energy_mix_list", "positive", "prices_costs"),   # T1b
                   ("S4", "substantive", "not_about_gas", "prices_costs"),  # T4
                   ("S5", "resource_list", "not_about_gas", "economic_opportunity")])  # T1
    p = lb.validate(bad)
    assert p.groupby("sample_id")["variable"].apply(list).to_dict() == {
        "S1": ["gas_mention"], "S2": ["gas_tone"], "S3": ["gas_tone"], "S4": ["gas_tone"], "S5": ["frame"]}


def test_write_xlsx_dropdowns_freeze_and_round_trip(tmp_path):
    from openpyxl import load_workbook

    sheet = lb.blind_sheet(lb.draw_sample(_corpus()))
    sheet.loc[0, "text"] = "Gas\u2019s \u201crise\u201d \u2013 " + "word " * 3000      # long, typographic
    path = tmp_path / "s.xlsx"
    lb.write_xlsx(sheet, path)
    wb = load_workbook(path)
    ws = wb["Labels"]
    assert ws.freeze_panes == "A2"
    assert [c.value for c in ws[1]] == ["sample_id", "title", "text", "gas_mention", "gas_tone", "frame",
                                        "unsure", "notes"]
    lists = {str(dv.sqref).split(":")[0][0]: dv.formula1 for dv in ws.data_validations.dataValidation}
    assert lists["D"] == '"substantive,resource_list,energy_mix_list,none"'
    assert lists["E"] == '"positive,neutral,negative,not_about_gas"'
    assert "resource_wealth" in lists["F"] and lists["G"] == '"x"'
    assert ws["C2"].alignment.wrap_text
    assert "Codebook" in wb.sheetnames
    with pytest.raises(FileExistsError):
        lb.write_xlsx(sheet, path)


def test_write_xlsx_refuses_text_over_cell_limit(tmp_path):
    sheet = lb.blind_sheet(lb.draw_sample(_corpus()))
    sheet.loc[3, "text"] = "x" * (lb.EXCEL_CELL_LIMIT + 1)
    with pytest.raises(ValueError, match=sheet.loc[3, "sample_id"]):
        lb.write_xlsx(sheet, tmp_path / "s.xlsx")


def test_read_sheet_from_xlsx(tmp_path):
    from openpyxl import load_workbook

    sheet = lb.blind_sheet(lb.draw_sample(_corpus()))
    path = tmp_path / "s.xlsx"
    lb.write_xlsx(sheet, path)
    wb = load_workbook(path)
    ws = wb["Labels"]
    ws["D2"], ws["E2"], ws["F2"], ws["G2"], ws["H2"] = "substantive", "positive", "prices_costs", "x", "check"
    wb.save(path)
    out = lb.read_sheet(path)
    first = out.iloc[0]
    assert (first["gas_mention"], first["gas_tone"], first["frame"], first["unsure"], first["notes"]) == \
        ("substantive", "positive", "prices_costs", True, "check")
    assert out["unsure"].sum() == 1 and (out["gas_mention"] == "").sum() == len(out) - 1


def _design_and_corpus():
    c = _corpus()
    return lb.design_table(lb.draw_sample(c)), c


def test_draw_replacement_same_stratum_unsampled_and_reproducible():
    design, c = _design_and_corpus()
    target = design[design["stratum"] == "2023_gas"].iloc[0]["sample_id"]
    out = lb.draw_replacement(design, c, target, "duplicate_of:S001")
    again = lb.draw_replacement(design, c, target, "duplicate_of:S001")
    pd.testing.assert_frame_equal(out, again)
    new = out.iloc[-1]
    assert new["sample_id"] == f"S{len(design) + 1:03d}" and new["replaces"] == target
    assert new["stratum"] == "2023_gas" and new["status"] == "active"
    assert new["post_id"] not in set(design["post_id"])
    old = out.set_index("sample_id").loc[target]
    assert (old["status"], old["reason"]) == ("excluded", "duplicate_of:S001")
    assert out["post_id"].is_unique


def test_second_replacement_uses_next_seed_and_never_reuses():
    design, c = _design_and_corpus()
    ids = design.loc[design["stratum"] == "2023_gas", "sample_id"].tolist()   # 40 articles, 20 sampled
    first = lb.draw_replacement(design, c, ids[0], "x")
    second = lb.draw_replacement(first, c, ids[1], "y")
    assert second["post_id"].is_unique and len(second) == len(design) + 2
    assert second.iloc[-1]["post_id"] != first.iloc[-1]["post_id"]
    with pytest.raises(ValueError):                       # already excluded
        lb.draw_replacement(second, c, ids[0], "again")


def test_draw_replacement_fails_when_stratum_exhausted():
    design, c = _design_and_corpus()
    target = design[design["stratum"] == "2024_other"].iloc[0]["sample_id"]   # all 9 already sampled
    with pytest.raises(ValueError, match="no unsampled"):
        lb.draw_replacement(design, c, target, "x")


def test_recode_draws_only_active_rows():
    design, c = _design_and_corpus()
    out = lb.draw_replacement(design, c, "S001", "x")
    assert "S001" not in lb.draw_recode(out, n=len(design), seed=1)


def test_append_to_xlsx_keeps_existing_cells_and_extends_dropdowns(tmp_path):
    from openpyxl import load_workbook

    sheet = lb.blind_sheet(lb.draw_sample(_corpus()))
    path = tmp_path / "s.xlsx"
    lb.write_xlsx(sheet, path)
    wb = load_workbook(path)
    wb["Labels"]["D2"], wb["Labels"]["H2"] = "substantive", "my note"
    wb.save(path)
    lb.append_to_xlsx(path, pd.DataFrame([{"sample_id": "S999", "title": "New", "text": "New text"}]))
    ws = load_workbook(path)["Labels"]
    last = ws.max_row
    assert last == len(sheet) + 2
    assert (ws["D2"].value, ws["H2"].value) == ("substantive", "my note")      # coder's work untouched
    assert [ws.cell(last, j).value for j in (1, 2, 3, 4)] == ["S999", "New", "New text", None]
    assert all(str(dv.sqref).endswith(str(last)) for dv in ws.data_validations.dataValidation)
    assert ws.cell(last, 3).alignment.wrap_text


def test_recompute_weights_after_corpus_shrinks_and_exclusion():
    design, c = _design_and_corpus()
    design = lb.draw_replacement(design, c, design[design["stratum"] == "2023_gas"].iloc[0]["sample_id"], "dup")
    smaller = c[c["post_id"] != c[(c["pub_date"].dt.year == 2023) & c["topics"].str.contains("natural_gas")]
                .iloc[-1]["post_id"]]                              # one 2023 gas article removed
    out = lb.recompute_weights(design, smaller)
    g = out[(out["stratum"] == "2023_gas") & (out["status"] == "active")]
    assert (g["N_h"] == 39).all() and (g["n_h"] == 20).all() and (g["weight"] == 1.95).all()
    assert out.loc[out["status"] == "excluded", "weight"].isna().all()
    assert out.loc[out["status"] == "active"].groupby("stratum")["weight"].sum().round(6).to_dict()["2022_gas"] == 25
