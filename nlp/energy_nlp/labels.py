"""Task 3: stratified hand-labelling sample, blind sheets, and label validation.

See nlp/labels/CODEBOOK.md. The sheet the coder sees holds only sample_id,
title and text; everything that could reveal the stratum (post id, date, URL,
weight) lives in sample_design.csv.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

CODES = {
    "gas_mention": ("substantive", "resource_list", "energy_mix_list", "none"),
    "gas_tone": ("positive", "neutral", "negative", "not_about_gas"),
    "frame": ("economic_opportunity", "supply_reliability", "prices_costs", "local_communities",
              "environment_climate", "governance_investment", "resource_wealth", "other"),
}
CODE_COLUMNS = list(CODES) + ["unsure", "notes"]
PER_STRATUM = {True: 20, False: 10}          # gas-tagged / other, per year
TRUTHY = {"1", "x", "y", "yes", "true", "t"}


def draw_sample(corpus: pd.DataFrame, seed: int = 42,
                per_stratum: dict[bool, int] = PER_STRATUM) -> pd.DataFrame:
    """Stratified sample (year x gas-tagged) with weights w = N_h / n_h.

    Rows come back in a random order and sample_id follows that order, so neither
    the id nor the row position reveals the stratum.
    """
    df = corpus.assign(year=corpus["pub_date"].dt.year,
                       gas_tagged=corpus["topics"].str.split(";").map(lambda t: "natural_gas" in t))
    rng = np.random.default_rng(seed)
    parts = []
    for (year, gas), g in df.sort_values("post_id").groupby(["year", "gas_tagged"]):
        n = min(per_stratum[gas], len(g))
        pick = g.iloc[np.sort(rng.choice(len(g), n, replace=False))]
        parts.append(pick.assign(stratum=f"{year}_{'gas' if gas else 'other'}", N_h=len(g), n_h=n))
    s = pd.concat(parts)
    s = s.iloc[rng.permutation(len(s))].reset_index(drop=True)
    s["weight"] = (s["N_h"] / s["n_h"]).round(4)
    s.insert(0, "sample_id", [f"S{i:03d}" for i in range(1, len(s) + 1)])
    return s


def design_table(sample: pd.DataFrame) -> pd.DataFrame:
    return sample[["sample_id", "post_id", "pub_date", "url", "year", "gas_tagged", "stratum",
                   "N_h", "n_h", "weight"]].assign(pub_date=lambda d: d["pub_date"].dt.date)


def blind_sheet(sample: pd.DataFrame) -> pd.DataFrame:
    """What the coder sees: id, title, text and empty code columns."""
    sheet = sample[["sample_id", "title", "text"]].copy()
    for c in CODE_COLUMNS:
        sheet[c] = ""
    return sheet


def draw_recode(design: pd.DataFrame, n: int = 30, seed: int = 43) -> list[str]:
    """sample_ids for the intra-coder reliability re-coding, in a fresh random order."""
    rng = np.random.default_rng(seed)
    ids = design["sample_id"].sort_values().to_numpy()
    return list(ids[rng.choice(len(ids), n, replace=False)])


def write_sheet(sheet: pd.DataFrame, path: Path, force: bool = False) -> None:
    """UTF-8 with BOM so Excel shows quotes and dashes correctly. Never overwrites
    an existing sheet (it may hold the coder's work) unless ``force``."""
    if path.exists() and not force:
        raise FileExistsError(f"{path} already exists - it may contain your labels. Use --force to replace it.")
    path.parent.mkdir(parents=True, exist_ok=True)
    sheet.to_csv(path, index=False, encoding="utf-8-sig")


EXCEL_CELL_LIMIT = 32_767
CODEBOOK_SUMMARY = [
    ("Order", "", "Code gas_mention first, then gas_tone, then frame. Natural gas = methane, CNG, LNG, gas "
                  "fields/pipelines/gas power. NOT LPG cylinders, crude oil/EACOP, petrol/diesel/HFO, hydropower."),
    ("gas_mention", "substantive", "At least one full sentence says something about gas itself (T4)."),
    ("gas_mention", "resource_list", "Gas only as an item in a list of national resources (T1)."),
    ("gas_mention", "energy_mix_list", "Gas only as an item in a list of energy sources (T1b)."),
    ("gas_mention", "none", "No natural gas mention."),
    ("gas_tone", "positive", "Benefit/success/progress, incl. factual growth ('production doubled') (T2a)."),
    ("gas_tone", "neutral", "Procedural, no outcome: tenders, meetings, bare announcements, levels without "
                            "change; gas replaced (e.g. by JNHPP) without unfavourable judgement (T5)."),
    ("gas_tone", "negative", "Problem/failure/risk/cost/delay, incl. factual decline (T2a) and unfavourable "
                             "comparison with another source (T5)."),
    ("gas_tone", "not_about_gas", "Always when gas_mention is not substantive."),
    ("frame", "economic_opportunity", "Growth, investment, jobs, industry, revenue as outcomes."),
    ("frame", "supply_reliability", "Generation, capacity, outages, supply volumes, infrastructure performance."),
    ("frame", "prices_costs", "Prices, tariffs, subsidies, costs or savings."),
    ("frame", "local_communities", "Host communities: benefits, grievances, compensation, CSR, local content."),
    ("frame", "environment_climate", "Emissions, deforestation, clean cooking, energy transition."),
    ("frame", "governance_investment", "Process: contracts, negotiations (HGA), licensing, regulation, SOEs."),
    ("frame", "resource_wealth", "Gas as part of national resource wealth; required when resource_list (T1)."),
    ("frame", "other", "None fits (accidents, courts, sport, obituaries)."),
    ("Rules", "T2", "Mixed tone: overall balance; if even, headline + first 3 paragraphs; if still even, neutral."),
    ("Rules", "T3", "Plans count as positive only if presented as beneficial; bare announcement = neutral."),
    ("Rules", "F1-F3", "One frame: headline/lead decides, then most text. Deals/regulation = governance; "
                       "outcomes = economic. Static asset in a list = resource_wealth."),
    ("unsure", "x", "Pick x if you hesitated; in notes write the variable and the two codes."),
    ("Full manual", "", "nlp/labels/CODEBOOK.md"),
]


def write_xlsx(sheet: pd.DataFrame, path: Path, force: bool = False) -> None:
    """Labelling workbook: dropdowns with only valid codes, an 'x' dropdown for
    unsure, wrapped text, frozen header and a codebook tab. Refuses to truncate
    text at Excel's cell limit and never overwrites an existing workbook unless
    ``force``. The written text is read back and compared with the input."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    if path.exists() and not force:
        raise FileExistsError(f"{path} already exists - it may contain your labels. Use --force to replace it.")
    too_long = sheet.loc[sheet["text"].str.len() > EXCEL_CELL_LIMIT, "sample_id"].tolist()
    if too_long:
        raise ValueError(f"text longer than Excel's {EXCEL_CELL_LIMIT}-character cell limit: {too_long}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Labels"
    cols = ["sample_id", "title", "text"] + CODE_COLUMNS
    ws.append(cols)
    for row in sheet[cols].itertuples(index=False):
        ws.append(list(row))
    n = len(sheet) + 1
    widths = {"sample_id": 9, "title": 28, "text": 110, "gas_mention": 17, "gas_tone": 15, "frame": 23,
              "unsure": 8, "notes": 30}
    head_fill = PatternFill("solid", fgColor="DCE6F2")
    code_fill = PatternFill("solid", fgColor="FFF8E1")
    for j, c in enumerate(cols, start=1):
        letter = ws.cell(1, j).column_letter
        ws.column_dimensions[letter].width = widths[c]
        ws.cell(1, j).font = Font(bold=True)
        ws.cell(1, j).fill = head_fill
        for i in range(2, n + 1):
            cell = ws.cell(i, j)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if c in CODE_COLUMNS:
                cell.fill = code_fill
    for i in range(2, n + 1):
        ws.row_dimensions[i].height = 409          # Excel's maximum; long texts scroll in the formula bar
    ws.freeze_panes = "A2"

    for c in list(CODES) + ["unsure"]:
        letter = ws.cell(1, cols.index(c) + 1).column_letter
        values = CODES[c] if c in CODES else ("x",)
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid code",
                            error="Pick a code from the list (see the Codebook tab).")
        dv.add(f"{letter}2:{letter}{n}")
        ws.add_data_validation(dv)

    cb = wb.create_sheet("Codebook")
    cb.append(["variable", "code", "meaning"])
    for r in CODEBOOK_SUMMARY:
        cb.append(list(r))
    for letter, w in zip("ABC", (14, 24, 110)):
        cb.column_dimensions[letter].width = w
    for row in cb.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in cb[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
    cb.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    back = [r[2] for r in load_workbook(path, read_only=True)["Labels"].iter_rows(min_row=2, values_only=True)]
    if back != sheet["text"].tolist():
        raise RuntimeError("article text changed when written to Excel")


def read_sheet(path: Path) -> pd.DataFrame:
    """Read a completed sheet: .xlsx (Labels tab), or CSV as saved by Excel
    (UTF-8, UTF-8 with BOM, or Windows-1252)."""
    if Path(path).suffix.lower() == ".xlsx":
        df = pd.read_excel(path, sheet_name="Labels", dtype=str, keep_default_na=False)
    else:
        for enc in ("utf-8-sig", "cp1252"):
            try:
                df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding=enc)
                break
            except UnicodeDecodeError:
                continue
    df.columns = [c.strip().lower() for c in df.columns]
    for c in CODES:
        df[c] = df[c].str.strip().str.lower().str.replace(" ", "_")
    df["unsure"] = df["unsure"].str.strip().str.lower().isin(TRUTHY)
    df["notes"] = df["notes"].str.strip()
    return df[["sample_id"] + CODE_COLUMNS]


def validate(labels: pd.DataFrame) -> pd.DataFrame:
    """Problems to fix by hand: missing or invalid codes and codebook consistency rules."""
    probs = []
    for _, r in labels.iterrows():
        for var, allowed in CODES.items():
            if r[var] == "":
                probs.append((r["sample_id"], var, "missing"))
            elif r[var] not in allowed:
                probs.append((r["sample_id"], var, f"invalid code {r[var]!r}"))
        if r["gas_mention"] in ("resource_list", "energy_mix_list", "none") \
                and r["gas_tone"] not in ("", "not_about_gas"):
            probs.append((r["sample_id"], "gas_tone",
                          f"gas_mention={r['gas_mention']} requires gas_tone=not_about_gas (T1/T1b/T4)"))
        if r["gas_mention"] == "substantive" and r["gas_tone"] == "not_about_gas":
            probs.append((r["sample_id"], "gas_tone",
                          "gas_mention=substantive but gas_tone=not_about_gas (T4)"))
        if r["gas_mention"] == "resource_list" and r["frame"] not in ("", "resource_wealth"):
            probs.append((r["sample_id"], "frame", "gas_mention=resource_list requires frame=resource_wealth (T1)"))
    return pd.DataFrame(probs, columns=["sample_id", "variable", "problem"])
